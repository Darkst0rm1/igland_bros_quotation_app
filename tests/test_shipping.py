"""Container shipping: sizes, freight methods, tier warnings and compatibility."""

from __future__ import annotations

import datetime as dt
from decimal import Decimal as D
from io import BytesIO

import pytest

from modules import (
    document_model,
    docx_generator,
    pdf_generator,
    pricing_service,
    quotation_service,
    revision_service,
    shipping_service,
)
from modules.authorization import PermissionDenied
from modules.capacity_importer import read_workbook
from modules.catalogue_service import create_product, create_variant, set_price
from modules.constants import (
    CHARGE_SOURCE_SHIPMENT,
    ChargeType,
    ContainerSize,
    ContainerType,
    FreightMethod,
    PriceTierCode,
    PriceWarningCode,
    RoleCode,
)
from modules.customer_service import create_customer
from modules.models import Quotation, QuotationCharge, ShippingLine
from modules.shipping_service import ShippingError
from modules.validation import CustomerInput, PriceInput, ProductInput, VariantInput

JAN = dt.date(2026, 1, 1)
QUOTE_DAY = dt.date(2026, 8, 4)


# --------------------------------------------------------------------------- #
# Fixtures
# --------------------------------------------------------------------------- #

@pytest.fixture
def admin(make_auth_user):
    return make_auth_user(RoleCode.SYS_ADMIN.value, username="root")


@pytest.fixture
def sales(make_auth_user):
    return make_auth_user(RoleCode.SALES.value, username="alice")


@pytest.fixture
def manager(make_auth_user):
    return make_auth_user(RoleCode.SALES_MANAGER.value, username="mgr")


@pytest.fixture
def carrier(session, admin):
    from seeds import seed_shipping

    seed_shipping.run(session)
    session.commit()
    return session.query(ShippingLine).order_by(ShippingLine.sort_order).first()


@pytest.fixture
def quotation(session, admin):
    """A two-line quotation at eight-container pricing."""
    customer = create_customer(
        session, admin,
        CustomerInput(customer_number="CUST-0001", company_name="Bunzl Canada"),
    )
    session.flush()

    quote = quotation_service.create_draft(
        session, admin, customer.id, quote_date=QUOTE_DAY
    )
    for size, price in (('12" White', "6.98"), ('14" White', "8.58")):
        product = create_product(
            session, admin,
            ProductInput(
                item_number=f"WB-{size[:2].strip()}", name=size, size_label=size,
                flute="B", depth_in=D("2"),
            ),
        )
        session.flush()
        variant = create_variant(
            session, admin, product.id,
            VariantInput(
                variant_item_number=f"WB-{size[:2].strip()}-A",
                board_quality="WT110 HPFL115 KM135", case_pack=50,
            ),
        )
        set_price(
            session, admin,
            PriceInput(
                product_variant_id=variant.id,
                price_tier_code=PriceTierCode.EIGHT_CONTAINER.value,
                price_per_pack=D(price), effective_from=JAN,
            ),
        )
        quotation_service.add_line(
            session, admin, quote,
            product_variant_id=variant.id,
            price_tier_code=PriceTierCode.EIGHT_CONTAINER.value,
            quantity_packs=D("1000"),
        )
    session.commit()
    return quote


def _add(session, user, quotation, carrier, size, count="1", freight="0"):
    return shipping_service.add_container(
        session, user, quotation,
        shipping_line_id=carrier.id,
        container_size=size,
        container_count=D(count),
        freight_cost=D(freight),
    )


# --------------------------------------------------------------------------- #
# Container sizes
# --------------------------------------------------------------------------- #

class TestContainerSizes:
    @pytest.mark.parametrize(
        ("size", "label"),
        [
            (ContainerSize.TWENTY_FT, "20 ft"),
            (ContainerSize.FORTY_FT, "40 ft"),
            (ContainerSize.FORTY_FT_HC, "40 ft High Cube"),
            (ContainerSize.FORTY_FIVE_FT_HC, "45 ft High Cube"),
        ],
    )
    def test_each_size_can_be_recorded(
        self, session, admin, quotation, carrier, size, label
    ):
        container = _add(session, admin, quotation, carrier, size)
        session.commit()
        assert container.container_size is size
        assert container.size_label == label
        assert container.container_type is ContainerType.DRY  # the default

    def test_a_custom_size_uses_the_free_text(self, session, admin, quotation, carrier):
        container = shipping_service.add_container(
            session, admin, quotation,
            shipping_line_id=carrier.id,
            container_size=ContainerSize.CUSTOM,
            custom_container_size="53 ft domestic",
            container_count=D("1"),
        )
        session.commit()
        assert container.size_label == "53 ft domestic"

    def test_the_default_is_forty_foot_high_cube_dry(
        self, session, admin, quotation, carrier
    ):
        """The price list ships in 40' HC containers, floor loaded."""
        container = shipping_service.add_container(
            session, admin, quotation, shipping_line_id=carrier.id
        )
        session.commit()
        assert container.container_size is ContainerSize.FORTY_FT_HC
        assert container.container_type is ContainerType.DRY

    def test_sizes_may_be_mixed_on_one_quotation(
        self, session, admin, quotation, carrier
    ):
        for size, count in (
            (ContainerSize.FORTY_FT_HC, "2"),
            (ContainerSize.TWENTY_FT, "1"),
            (ContainerSize.FORTY_FT, "1"),
        ):
            _add(session, admin, quotation, carrier, size, count)
        session.commit()
        assert shipping_service.total_containers(session, quotation.id) == D("4")

    def test_different_carriers_per_container(self, session, admin, quotation, carrier):
        other = shipping_service.create_shipping_line(session, admin, "Custom Carrier")
        session.flush()
        first = _add(session, admin, quotation, carrier, ContainerSize.FORTY_FT_HC)
        second = shipping_service.add_container(
            session, admin, quotation, shipping_line_id=other.id
        )
        session.commit()
        assert first.carrier_name != second.carrier_name

    def test_a_carrier_off_the_list_uses_free_text(self, session, admin, quotation):
        container = shipping_service.add_container(
            session, admin, quotation, custom_shipping_line="Regional Feeder Line"
        )
        session.commit()
        assert container.carrier_name == "Regional Feeder Line"

    def test_a_zero_count_is_refused(self, session, admin, quotation, carrier):
        with pytest.raises(ShippingError, match="at least one"):
            _add(session, admin, quotation, carrier, ContainerSize.FORTY_FT, count="0")

    def test_arrival_before_departure_is_refused(
        self, session, admin, quotation, carrier
    ):
        with pytest.raises(ShippingError, match="before the departure"):
            shipping_service.add_container(
                session, admin, quotation, shipping_line_id=carrier.id,
                estimated_departure=dt.date(2026, 9, 10),
                estimated_arrival=dt.date(2026, 9, 1),
            )

    def test_transit_days_are_derived_from_the_dates(
        self, session, admin, quotation, carrier
    ):
        container = shipping_service.add_container(
            session, admin, quotation, shipping_line_id=carrier.id,
            estimated_departure=dt.date(2026, 9, 1),
            estimated_arrival=dt.date(2026, 9, 29),
        )
        session.commit()
        assert container.transit_days == 28

    def test_arrival_is_derived_from_transit_days(
        self, session, admin, quotation, carrier
    ):
        container = shipping_service.add_container(
            session, admin, quotation, shipping_line_id=carrier.id,
            estimated_departure=dt.date(2026, 9, 1), transit_days=28,
        )
        session.commit()
        assert container.estimated_arrival == dt.date(2026, 9, 29)


# --------------------------------------------------------------------------- #
# Pricing-tier warnings
# --------------------------------------------------------------------------- #

class TestTierWarnings:
    @staticmethod
    def _short_warnings(session, quotation):
        return [
            w for w in pricing_service.evaluate_quotation(session, quotation)
            if w.code is PriceWarningCode.TIER_CONTAINERS_SHORT
        ]

    def test_eight_container_pricing_warns_below_eight(
        self, session, admin, quotation, carrier
    ):
        _add(session, admin, quotation, carrier, ContainerSize.FORTY_FT_HC, count="4")
        session.commit()

        warnings = self._short_warnings(session, quotation)
        assert len(warnings) == 1
        assert "4" in warnings[0].message and "8" in warnings[0].message

    def test_eight_container_pricing_is_satisfied_at_eight(
        self, session, admin, quotation, carrier
    ):
        _add(session, admin, quotation, carrier, ContainerSize.FORTY_FT_HC, count="8")
        session.commit()
        assert self._short_warnings(session, quotation) == []

    def test_three_container_pricing_warns_below_three(
        self, session, admin, manager, quotation, carrier
    ):
        for item in quotation.items:
            quotation_service.change_line_tier(
                session, admin, quotation, item.id,
                PriceTierCode.THREE_CONTAINER.value,
            ) if False else None
        # Re-price both lines onto the three-container tier.
        from modules.catalogue_service import set_price as add_price
        from modules.repositories import get_variant

        for item in quotation.items:
            variant = get_variant(session, item.product_variant_id)
            add_price(
                session, admin,
                PriceInput(
                    product_variant_id=variant.id,
                    price_tier_code=PriceTierCode.THREE_CONTAINER.value,
                    price_per_pack=D("7.20"), effective_from=JAN,
                ),
            )
            quotation_service.change_line_tier(
                session, admin, quotation, item.id,
                PriceTierCode.THREE_CONTAINER.value,
            )
        _add(session, admin, quotation, carrier, ContainerSize.TWENTY_FT, count="2")
        session.commit()

        warnings = self._short_warnings(session, quotation)
        assert len(warnings) == 1
        assert "Three Containers" in warnings[0].message

    def test_standard_pricing_never_warns_about_containers(
        self, session, admin, quotation, carrier
    ):
        from modules.catalogue_service import set_price as add_price
        from modules.repositories import get_variant

        for item in quotation.items:
            variant = get_variant(session, item.product_variant_id)
            add_price(
                session, admin,
                PriceInput(
                    product_variant_id=variant.id,
                    price_tier_code=PriceTierCode.STANDARD.value,
                    price_per_pack=D("7.42"), effective_from=JAN,
                ),
            )
            quotation_service.change_line_tier(
                session, admin, quotation, item.id, PriceTierCode.STANDARD.value
            )
        _add(session, admin, quotation, carrier, ContainerSize.TWENTY_FT, count="1")
        session.commit()
        assert self._short_warnings(session, quotation) == []

    def test_the_warning_says_the_tier_was_not_changed(
        self, session, admin, quotation, carrier
    ):
        """The obvious next question is 'did it just reprice my quotation?'."""
        _add(session, admin, quotation, carrier, ContainerSize.FORTY_FT, count="1")
        session.commit()

        warning = self._short_warnings(session, quotation)[0]
        assert "not been changed" in warning.message
        assert quotation.items[0].price_per_pack == D("6.98")

    def test_container_rows_take_precedence_over_the_line_field(
        self, session, admin, quotation, carrier
    ):
        """The shipment is the real plan once one exists."""
        quotation_service.update_line(
            session, admin, quotation, quotation.items[0].id,
            container_count=D("9"),
        )
        session.commit()
        assert self._short_warnings(session, quotation) == []

        _add(session, admin, quotation, carrier, ContainerSize.FORTY_FT_HC, count="2")
        session.commit()
        # 2 from the shipment, not 9 from the line.
        assert len(self._short_warnings(session, quotation)) == 1

    def test_a_quotation_with_no_shipment_uses_the_line_field(
        self, session, admin, quotation
    ):
        """Quotations raised before container shipping existed keep working."""
        quotation_service.update_line(
            session, admin, quotation, quotation.items[0].id,
            container_count=D("9"),
        )
        session.commit()
        assert shipping_service.get_shipment(session, quotation.id) is None
        assert self._short_warnings(session, quotation) == []


# --------------------------------------------------------------------------- #
# Freight
# --------------------------------------------------------------------------- #

class TestFreight:
    @pytest.fixture
    def shipped(self, session, admin, quotation, carrier):
        _add(session, admin, quotation, carrier, ContainerSize.FORTY_FT_HC,
             count="2", freight="3200")
        _add(session, admin, quotation, carrier, ContainerSize.TWENTY_FT,
             count="1", freight="1900")
        session.commit()
        return quotation

    @staticmethod
    def _charges(session, quotation_id):
        return session.query(QuotationCharge).filter_by(quotation_id=quotation_id).all()

    def test_total_freight_multiplies_by_the_container_count(self, session, shipped):
        shipment = shipping_service.get_shipment(session, shipped.id)
        assert shipment.total_freight == D("8300.00")  # 2x3200 + 1x1900

    def test_freight_per_container(self, session, shipped):
        assert shipping_service.freight_per_container(session, shipped.id) == D("2766.67")

    def test_included_freight_creates_no_charge(self, session, admin, shipped):
        shipping_service.update_shipment(
            session, admin, shipped, freight_method=FreightMethod.INCLUDED
        )
        session.commit()
        assert self._charges(session, shipped.id) == []
        assert shipped.charges_total == D("0")

    def test_internal_only_freight_creates_no_charge(self, session, admin, shipped):
        """An internal-only *charge* would still be added to the grand total, so
        internal freight must not be a charge at all."""
        shipping_service.update_shipment(
            session, admin, shipped, freight_method=FreightMethod.INTERNAL_ONLY
        )
        session.commit()
        assert self._charges(session, shipped.id) == []
        assert shipped.grand_total == shipped.subtotal

    def test_separate_freight_creates_exactly_one_charge(self, session, admin, shipped):
        shipping_service.update_shipment(
            session, admin, shipped, freight_method=FreightMethod.ADDED_SEPARATELY
        )
        session.commit()

        charges = self._charges(session, shipped.id)
        assert len(charges) == 1
        assert charges[0].source == CHARGE_SOURCE_SHIPMENT
        assert charges[0].amount == D("8300.00")
        assert charges[0].is_customer_visible
        assert shipped.grand_total == shipped.subtotal + D("8300.00")

    def test_repeated_syncs_do_not_duplicate_the_charge(self, session, admin, shipped):
        shipping_service.update_shipment(
            session, admin, shipped, freight_method=FreightMethod.ADDED_SEPARATELY
        )
        for _ in range(5):
            shipping_service.sync_freight(session, admin, shipped)
        session.commit()

        assert len(self._charges(session, shipped.id)) == 1
        assert shipped.charges_total == D("8300.00")

    def test_switching_method_removes_the_charge_again(self, session, admin, shipped):
        shipping_service.update_shipment(
            session, admin, shipped, freight_method=FreightMethod.ADDED_SEPARATELY
        )
        session.commit()
        assert len(self._charges(session, shipped.id)) == 1

        shipping_service.update_shipment(
            session, admin, shipped, freight_method=FreightMethod.INCLUDED
        )
        session.commit()
        assert self._charges(session, shipped.id) == []

    def test_adding_a_container_updates_the_charge(self, session, admin, shipped, carrier):
        shipping_service.update_shipment(
            session, admin, shipped, freight_method=FreightMethod.ADDED_SEPARATELY
        )
        session.commit()

        _add(session, admin, shipped, carrier, ContainerSize.FORTY_FT,
             count="1", freight="2800")
        session.commit()

        charges = self._charges(session, shipped.id)
        assert len(charges) == 1
        assert charges[0].amount == D("11100.00")

    def test_a_manual_freight_charge_alongside_a_shipment_warns(
        self, session, admin, shipped
    ):
        quotation_service.add_charge(
            session, admin, shipped, charge_type=ChargeType.FREIGHT,
            description="Inland haulage", quantity=D("1"), rate=D("500"),
        )
        session.commit()

        codes = {
            w.code for w in pricing_service.evaluate_quotation(session, shipped)
        }
        assert PriceWarningCode.DUPLICATE_FREIGHT in codes

    def test_no_duplicate_warning_without_a_manual_charge(self, session, shipped):
        codes = {
            w.code for w in pricing_service.evaluate_quotation(session, shipped)
        }
        assert PriceWarningCode.DUPLICATE_FREIGHT not in codes

    def test_landed_freight_excludes_separately_charged_freight(
        self, session, admin, shipped
    ):
        """Charging the customer and counting it as cost would understate margin."""
        shipping_service.update_shipment(
            session, admin, shipped, freight_method=FreightMethod.INTERNAL_ONLY
        )
        session.commit()
        assert shipping_service.landed_freight(session, admin, shipped.id) == D("8300.00")

        shipping_service.update_shipment(
            session, admin, shipped, freight_method=FreightMethod.ADDED_SEPARATELY
        )
        session.commit()
        assert shipping_service.landed_freight(session, admin, shipped.id) is None

    def test_freight_is_apportioned_across_products(
        self, session, admin, shipped, carrier
    ):
        container = shipping_service.get_shipment(session, shipped.id).containers[0]
        for item in shipped.items:
            shipping_service.allocate_product(
                session, admin, shipped, container.id, item.id,
                quantity_per_container=D("500"),
            )
        session.commit()
        shipping_service.apportion_freight(session, shipped.id)
        session.commit()

        shares = [a.allocated_freight for a in container.allocations]
        assert len(shares) == 2
        # 2 containers x 3200, split evenly between two equal allocations.
        assert sum(shares) == D("6400.00")


class TestFreightPermissions:
    def test_sales_may_build_a_shipment(self, session, sales, admin, quotation, carrier):
        quotation.sales_user_id = sales.id
        session.commit()
        container = shipping_service.add_container(
            session, sales, quotation, shipping_line_id=carrier.id
        )
        session.commit()
        assert container.id

    def test_sales_may_not_enter_freight(self, session, sales, quotation, carrier):
        quotation.sales_user_id = sales.id
        session.commit()
        with pytest.raises(PermissionDenied):
            shipping_service.add_container(
                session, sales, quotation, shipping_line_id=carrier.id,
                freight_cost=D("3200"),
            )

    def test_sales_may_not_change_the_freight_method(self, session, sales, quotation):
        quotation.sales_user_id = sales.id
        session.commit()
        with pytest.raises(PermissionDenied):
            shipping_service.update_shipment(
                session, sales, quotation, freight_method=FreightMethod.ADDED_SEPARATELY
            )

    def test_a_manager_may_do_both(self, session, manager, quotation, carrier):
        quotation.sales_user_id = manager.id
        session.commit()
        shipping_service.add_container(
            session, manager, quotation, shipping_line_id=carrier.id,
            freight_cost=D("3200"),
        )
        shipping_service.update_shipment(
            session, manager, quotation, freight_method=FreightMethod.ADDED_SEPARATELY
        )
        session.commit()
        assert shipping_service.get_shipment(session, quotation.id).total_freight == D("3200.00")

    def test_reading_freight_needs_the_permission(self, session, sales, quotation):
        with pytest.raises(PermissionDenied):
            shipping_service.landed_freight(session, sales, quotation.id)


# --------------------------------------------------------------------------- #
# Capacity import
# --------------------------------------------------------------------------- #

class TestCapacityImport:
    @staticmethod
    def _workbook(rows):
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.append(["WHITE BOXES B FLUTE"])          # title bar, row 1
        ws.append([])                               # row 2
        ws.append(["Product", "Bundles Per Container"])  # header on row 3
        for label, bundles in rows:
            ws.append([label, bundles])
        ws.append(["Container type: 40' HC."])
        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer

    def test_the_header_is_found_below_the_title_bar(self):
        plan = read_workbook(self._workbook([('7" White', 3500), ('8" White', 2940)]))
        assert plan.header_row == 3
        assert len(plan.rows) == 2

    def test_the_container_is_read_from_the_sheet_note(self):
        plan = read_workbook(self._workbook([('7" White', 3500)]))
        assert plan.container_size is ContainerSize.FORTY_FT_HC

    def test_capacity_falling_with_size_is_not_flagged(self):
        plan = read_workbook(
            self._workbook([('7" White', 3500), ('12" White', 1880), ('18" White', 890)])
        )
        assert plan.counts()["anomalous"] == 0

    def test_capacity_rising_with_size_is_flagged(self):
        """A bigger box fitting more per container contradicts itself."""
        plan = read_workbook(
            self._workbook(
                [('7" White', 3500), ('18" White', 890), ('20" White', 1512)]
            )
        )
        assert plan.counts()["anomalous"] == 1
        flagged = next(r for r in plan.rows if r.is_anomalous)
        assert flagged.product_label == '20" White'
        assert "1,512" in flagged.anomaly_note
        assert flagged.bundles_per_container == D("1512")  # imported as given

    def test_the_import_writes_capacity_against_products(
        self, session, admin, quotation
    ):
        from modules.capacity_importer import commit
        from modules.repositories import find_product_by_size

        plan = read_workbook(self._workbook([('12" White', 1880)]))
        summary = commit(session, admin, plan, "bundles.xlsx")
        session.commit()

        assert summary["created"] == 1
        product = find_product_by_size(session, '12" White')
        capacity = shipping_service.container_capacity(
            session, product.id, ContainerSize.FORTY_FT_HC, ContainerType.DRY
        )
        assert capacity.bundles_per_container == D("1880")
        # The source does not say what a bundle holds, so pieces stay unknown.
        assert capacity.units_per_bundle is None
        assert capacity.pieces_per_container is None

    def test_a_product_that_does_not_exist_is_skipped_not_invented(
        self, session, admin, quotation
    ):
        from modules.capacity_importer import commit

        plan = read_workbook(self._workbook([('99" White', 100)]))
        summary = commit(session, admin, plan, "bundles.xlsx")
        session.commit()
        assert summary["skipped"] == 1 and summary["created"] == 0


class TestAllocation:
    def test_quantities_derive_from_recorded_capacity(
        self, session, admin, quotation, carrier
    ):
        from modules.capacity_importer import commit

        plan = read_workbook(TestCapacityImport._workbook([('12" White', 1880)]))
        commit(session, admin, plan, "bundles.xlsx")
        session.commit()

        container = _add(
            session, admin, quotation, carrier, ContainerSize.FORTY_FT_HC, count="3"
        )
        line = quotation.items[0]
        allocation = shipping_service.allocate_product(
            session, admin, quotation, container.id, line.id
        )
        session.commit()

        assert allocation.quantity_per_container == D("1880")
        assert allocation.total_allocated_quantity == D("5640")  # 1880 x 3

    def test_without_capacity_the_quantity_must_be_supplied(
        self, session, admin, quotation, carrier
    ):
        container = _add(session, admin, quotation, carrier, ContainerSize.TWENTY_FT)
        with pytest.raises(ShippingError, match="cannot be derived"):
            shipping_service.allocate_product(
                session, admin, quotation, container.id, quotation.items[0].id
            )

    def test_changing_the_container_count_updates_the_total(
        self, session, admin, quotation, carrier
    ):
        container = _add(session, admin, quotation, carrier, ContainerSize.FORTY_FT, count="2")
        allocation = shipping_service.allocate_product(
            session, admin, quotation, container.id, quotation.items[0].id,
            quantity_per_container=D("1000"),
        )
        session.commit()
        assert allocation.total_allocated_quantity == D("2000")

        shipping_service.update_container(
            session, admin, quotation, container.id, container_count=D("5")
        )
        session.commit()
        assert allocation.total_allocated_quantity == D("5000")


# --------------------------------------------------------------------------- #
# Documents and revisions
# --------------------------------------------------------------------------- #

class TestDocuments:
    @pytest.fixture
    def shipped(self, session, admin, quotation, carrier):
        _add(session, admin, quotation, carrier, ContainerSize.FORTY_FT_HC,
             count="2", freight="3200")
        shipping_service.update_shipment(
            session, admin, quotation,
            port_of_loading="Çerkezköy", port_of_discharge="Toronto",
            freight_method=FreightMethod.INTERNAL_ONLY,
        )
        session.commit()
        return quotation

    def test_the_section_is_absent_unless_enabled(self, session, shipped):
        model = document_model.build_document(session, shipped)
        assert model.shipping is None

    def test_enabling_it_adds_the_table(self, session, admin, shipped):
        shipping_service.update_shipment(
            session, admin, shipped, show_on_document=True
        )
        session.commit()

        model = document_model.build_document(session, shipped)
        assert model.shipping is not None
        assert model.shipping.rows[0][1] == "40 ft High Cube"

    def test_internal_freight_never_reaches_the_document(
        self, session, admin, shipped
    ):
        from pypdf import PdfReader

        shipping_service.update_shipment(
            session, admin, shipped, show_on_document=True,
            customer_visible_freight=False,
        )
        session.commit()

        model = document_model.build_document(session, shipped)
        assert "freight" not in model.shipping.columns
        assert "6,400" not in repr(model)

        text = PdfReader(BytesIO(pdf_generator.render(model))).pages[0].extract_text()
        assert "6,400" not in text
        assert "40 ft High Cube" in text

        from docx import Document as DocxDocument

        document = DocxDocument(BytesIO(docx_generator.render(model)))
        cells = " ".join(
            c.text for t in document.tables for r in t.rows for c in r.cells
        )
        assert "6,400" not in cells
        assert "40 ft High Cube" in cells

    def test_freight_appears_only_when_marked_visible(self, session, admin, shipped):
        shipping_service.update_shipment(
            session, admin, shipped, show_on_document=True,
            customer_visible_freight=True,
        )
        session.commit()
        model = document_model.build_document(session, shipped)
        assert "freight" in model.shipping.columns

    def test_a_quotation_without_shipping_is_unaffected(
        self, session, admin, quotation
    ):
        model = document_model.build_document(session, quotation)
        assert model.shipping is None
        assert pdf_generator.render(model).startswith(b"%PDF-")


class TestRevisions:
    def test_a_revision_carries_the_shipping_plan(
        self, session, admin, manager, quotation, carrier
    ):
        from modules import approval_service, document_service

        _add(session, admin, quotation, carrier, ContainerSize.FORTY_FT_HC,
             count="2", freight="3200")
        _add(session, admin, quotation, carrier, ContainerSize.TWENTY_FT,
             count="1", freight="1900")
        session.commit()

        approval_service.submit(session, quotation, admin)
        session.commit()
        document_service.generate(
            session, admin, quotation, document_service.DocumentFormat.PDF, draft=False
        )
        revision_service.issue(session, admin, quotation)
        session.commit()

        revised = revision_service.create_revision(
            session, admin, quotation, "customer added a size"
        )
        session.commit()

        shipment = shipping_service.get_shipment(session, revised.id)
        assert shipment is not None
        assert len(shipment.containers) == 2
        assert shipment.total_freight == D("8300.00")

    def test_the_snapshot_records_the_containers(
        self, session, admin, quotation, carrier
    ):
        _add(session, admin, quotation, carrier, ContainerSize.FORTY_FT_HC, count="2")
        session.commit()

        snapshot = revision_service.snapshot(quotation)
        assert snapshot["shipment"]["containers"][0]["container_size"] == "40 ft High Cube"

    def test_a_snapshot_without_shipping_is_still_valid(self, session, quotation):
        """Snapshots taken before container shipping existed lack the key."""
        snapshot = revision_service.snapshot(quotation)
        assert snapshot["shipment"] is None

        older = dict(snapshot)
        older.pop("shipment")
        diff = revision_service.compare(older, snapshot)
        assert not revision_service.has_changes(diff)


class TestShippingReports:
    @pytest.fixture
    def shipped(self, session, admin, quotation, carrier):
        _add(session, admin, quotation, carrier, ContainerSize.FORTY_FT_HC,
             count="2", freight="3200")
        _add(session, admin, quotation, carrier, ContainerSize.TWENTY_FT,
             count="1", freight="1900")
        shipping_service.update_shipment(
            session, admin, quotation,
            port_of_loading="Çerkezköy", port_of_discharge="Toronto",
        )
        session.commit()
        return quotation

    def test_headlines(self, session, admin, shipped):
        from modules import reporting_service

        figures = reporting_service.shipping_headlines(session, admin)
        assert figures.total_containers == D("3")
        assert figures.total_freight == D("8300.00")
        assert figures.average_freight_per_container == D("2766.67")
        assert figures.quotations_with_shipping == 1

    def test_averages_are_none_rather_than_zero_when_unknown(self, session, admin):
        """Nothing shipped is not the same as an average of zero."""
        from modules import reporting_service

        figures = reporting_service.shipping_headlines(session, admin)
        assert figures.total_containers == D("0")
        assert figures.average_freight_per_container is None
        assert figures.average_transit_days is None

    def test_containers_by_size_uses_readable_labels(self, session, admin, shipped):
        from modules import reporting_service

        frame = reporting_service.containers_by_size(session, admin)
        sizes = set(frame["Container size"])
        assert sizes == {"40 ft High Cube", "20 ft"}

    def test_containers_by_shipping_line(self, session, admin, shipped, carrier):
        from modules import reporting_service

        frame = reporting_service.containers_by_shipping_line(session, admin)
        assert frame.iloc[0]["Shipping line"] == carrier.name
        assert frame.iloc[0]["Containers"] == 3.0

    def test_the_shipment_report_has_one_row_per_container(
        self, session, admin, shipped
    ):
        from modules import reporting_service

        frame = reporting_service.shipments(session, admin)
        assert len(frame) == 2
        assert set(frame["Port of discharge"]) == {"Toronto"}

    def test_every_shipping_aggregate_survives_an_empty_database(self, session, admin):
        from modules import reporting_service

        for builder in (
            reporting_service.containers_by_size,
            reporting_service.containers_by_type,
            reporting_service.containers_by_shipping_line,
            reporting_service.containers_by_route,
            reporting_service.shipments,
        ):
            frame = builder(session, admin)
            assert frame.empty
            assert len(frame.columns) >= 3, builder.__name__

    def test_filtering_by_container_size(self, session, admin, shipped):
        from modules import reporting_service
        from modules.reporting_service import ReportFilters

        matching = reporting_service.headlines(
            session, admin, ReportFilters(container_sizes=("20FT",))
        )
        assert matching.total == 1

        other = reporting_service.headlines(
            session, admin, ReportFilters(container_sizes=("45FT_HC",))
        )
        assert other.total == 0

    def test_filtering_by_carrier_and_port(self, session, admin, shipped, carrier):
        from modules import reporting_service
        from modules.reporting_service import ReportFilters

        assert reporting_service.headlines(
            session, admin, ReportFilters(shipping_line_ids=(carrier.id,))
        ).total == 1
        assert reporting_service.headlines(
            session, admin, ReportFilters(port_of_discharge="Toronto")
        ).total == 1
        assert reporting_service.headlines(
            session, admin, ReportFilters(port_of_discharge="Rotterdam")
        ).total == 0

    def test_a_multi_container_quotation_is_not_counted_twice(
        self, session, admin, shipped, carrier
    ):
        """Filtering joins through containers, so the quotation must not multiply."""
        from modules import reporting_service
        from modules.reporting_service import ReportFilters

        figures = reporting_service.headlines(
            session, admin, ReportFilters(shipping_line_ids=(carrier.id,))
        )
        assert figures.total == 1
        assert figures.total_quoted == shipped.grand_total

    def test_minimum_container_filter(self, session, admin, shipped):
        from modules import reporting_service
        from modules.reporting_service import ReportFilters

        assert reporting_service.headlines(
            session, admin, ReportFilters(min_containers=D("3"))
        ).total == 1
        assert reporting_service.headlines(
            session, admin, ReportFilters(min_containers=D("4"))
        ).total == 0


class TestShippingLineAdministration:
    def test_a_carrier_can_be_added_and_renamed(self, session, admin):
        line = shipping_service.create_shipping_line(session, admin, "Wan Hai")
        session.commit()
        assert line.name == "Wan Hai"

        shipping_service.update_shipping_line(
            session, admin, line.id, name="Wan Hai Lines",
            is_active=True, sort_order=50,
        )
        session.commit()
        assert line.name == "Wan Hai Lines"

    def test_a_duplicate_name_is_refused(self, session, admin, carrier):
        with pytest.raises(ShippingError, match="already on the list"):
            shipping_service.create_shipping_line(session, admin, carrier.name.lower())

    def test_a_blank_name_is_refused(self, session, admin):
        with pytest.raises(ShippingError, match="needs a name"):
            shipping_service.create_shipping_line(session, admin, "   ")

    def test_removal_is_a_soft_delete_that_preserves_history(
        self, session, admin, quotation, carrier
    ):
        """A carrier removed from the list must not blank out past quotations."""
        container = _add(
            session, admin, quotation, carrier, ContainerSize.FORTY_FT_HC
        )
        session.commit()

        shipping_service.delete_shipping_line(session, admin, carrier.id)
        session.commit()

        assert carrier.name not in [
            line.name for line in shipping_service.shipping_lines(session)
        ]
        assert container.carrier_name == carrier.name

    def test_sales_cannot_manage_carriers(self, session, sales):
        with pytest.raises(PermissionDenied):
            shipping_service.create_shipping_line(session, sales, "Anything")


class TestImmutability:
    def test_shipping_cannot_be_edited_once_issued(
        self, session, admin, quotation, carrier
    ):
        _add(session, admin, quotation, carrier, ContainerSize.FORTY_FT_HC)
        session.commit()

        quotation.is_locked = True
        session.commit()

        with pytest.raises(PermissionDenied, match="Create a revision"):
            shipping_service.update_shipment(
                session, admin, quotation, port_of_loading="Somewhere else"
            )
